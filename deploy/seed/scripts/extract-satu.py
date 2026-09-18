#!/usr/bin/env python3
"""Читает Satu.kz-экспорт (.xlsx) и пишет /tmp/satu-import.json для import-satu.js"""
import sys, json, re, html
import openpyxl

SRC = sys.argv[1] if len(sys.argv) > 1 else '/Users/mira/Downloads/export-products-28-08-26_10-04-07.xlsx'
OUT = '/tmp/satu-import.json'

wb = openpyxl.load_workbook(SRC, data_only=True)
prod_ws = wb['Export Products Sheet']
grp_ws = wb['Export Groups Sheet']

def strip_html(s):
    if not s:
        return ''
    s = re.sub(r'<br\s*/?>', '\n', s, flags=re.I)
    s = re.sub(r'</p>', '\n', s, flags=re.I)
    s = re.sub(r'<[^>]+>', '', s)
    s = html.unescape(s)
    s = re.sub(r'\n{3,}', '\n\n', s)
    return s.strip()

BRAND_FIX = {
    'hoffman': 'HOFMANN', 'hofmann': 'HOFMANN', 'HOFMANN': 'HOFMANN',
    'karcher': 'Kärcher', 'chance': 'Chancee',
    'ipc portotecnica': 'IPC Portotecnica', 'portotecnica': 'Portotecnica',
    'gansow': 'GANSOW', 'soteco': 'SOTECO', 'idrobase': 'IDROBASE',
    'samoa': 'SAMOA', 'ravaglioli': 'RAVAGLIOLI', 'elsea': 'Elsea',
    'nussbaum': 'Nussbaum', 'chicago pneumatic': 'Chicago Pneumatic',
    'электромотор': 'Электромотор',
}
def fix_brand(b):
    if not b:
        return ''
    return BRAND_FIX.get(str(b).strip().lower(), str(b).strip())

# ---------- ГРУППЫ ----------
grows = list(grp_ws.iter_rows(values_only=True))[1:]
groups = {}  # num -> {name, parent}
for r in grows:
    num, name, ident, parent = r[0], r[1], r[2], r[3]
    if num is None:
        continue
    groups[str(num)] = {
        'name': (name or '').strip(),
        'parent': str(parent) if parent else None,
        'image': r[9] or None,
    }

def root_of(num):
    seen = set()
    while num in groups and groups[num]['parent'] and groups[num]['parent'] in groups and num not in seen:
        seen.add(num)
        num = groups[num]['parent']
    return num

roots = [n for n, g in groups.items()
         if not g['parent'] or g['parent'] not in groups]

categories = []          # {num, name, image}
subcategories = []        # {num, name, category_num}
cat_by_num, sub_by_num = {}, {}
for n in roots:
    categories.append({'num': n, 'name': groups[n]['name'], 'image': groups[n]['image']})
    cat_by_num[n] = n
for n, g in groups.items():
    if n in cat_by_num:
        continue
    r = root_of(n)
    subcategories.append({'num': n, 'name': g['name'], 'category_num': r})
    sub_by_num[n] = r

# ---------- ТОВАРЫ ----------
prows = list(prod_ws.iter_rows(values_only=True))
hdr = prows[0]
H = {name: i for i, name in enumerate(hdr) if name and name not in ('Название_Характеристики', 'Измерение_Характеристики', 'Значение_Характеристики')}
SPEC0 = 39   # первый триплет характеристик

products = []
for r in prows[1:]:
    if not r or not r[H['Название_позиции']]:
        continue
    grp = str(r[H['Номер_группы']]) if r[H['Номер_группы']] else None
    if grp in cat_by_num:
        cat_num, sub_num = grp, None
    elif grp in sub_by_num:
        cat_num, sub_num = sub_by_num[grp], grp
    else:
        cat_num, sub_num = (roots[0] if roots else None), None

    # характеристики
    specs = {}
    i = SPEC0
    while i + 2 < len(r):
        nm, meas, val = r[i], r[i + 1], r[i + 2]
        if nm and val not in (None, ''):
            v = str(val).strip()
            if meas and str(meas).strip() and str(meas).strip().lower() not in v.lower():
                v = f'{v} {str(meas).strip()}'
            specs[str(nm).strip()] = v
        i += 3

    imgs_raw = r[H['Ссылка_изображения']] or ''
    images = [u.strip() for u in re.split(r'[,\s]+', str(imgs_raw)) if u.strip().startswith('http')]

    price_val = r[H['Цена']]
    try:
        price = float(price_val) if price_val not in (None, '') else 0
    except (TypeError, ValueError):
        price = 0

    products.append({
        'article': str(r[H['Код_товара']] or '').strip(),
        'name': str(r[H['Название_позиции']]).strip(),
        'description': strip_html(r[H['Описание']]),
        'search': str(r[H['Поисковые_запросы']] or '').strip()[:400],
        'price': price,
        'price_on_request': 1 if price <= 0 else 0,
        'in_stock': 1 if str(r[H['Наличие']] or '').strip() == '+' else 0,
        'brand': fix_brand(r[H['Производитель']]),
        'country': str(r[H['Страна_производитель']] or '').strip(),
        'category_num': cat_num,
        'subcategory_num': sub_num,
        'group_name': str(r[H['Название_группы']] or '').strip(),
        'specs': specs,
        'image_urls': images,
        'unit': str(r[H['Единица_измерения']] or 'шт').strip(),
    })

out = {'categories': categories, 'subcategories': subcategories, 'products': products}
json.dump(out, open(OUT, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
print(f'✅ {OUT}: {len(categories)} категорий, {len(subcategories)} подкатегорий, {len(products)} товаров')
print('   товаров с картинками:', sum(1 for p in products if p['image_urls']))
print('   всего URL картинок  :', sum(len(p['image_urls']) for p in products))
