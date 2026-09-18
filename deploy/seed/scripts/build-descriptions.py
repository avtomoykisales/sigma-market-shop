#!/usr/bin/env python3
# ---------------------------------------------------------------------------
# Достаёт описания категорий и подкатегорий из экспорта Satu.kz
# («Export Groups Sheet») и складывает их в deploy/seed/category-descriptions.json
# (ключ = точное «Название_группы», значение = очищенный HTML).
#
# Применяются они скриптом  scripts/import-descriptions.js  (по совпадению имени,
# только в пустое поле `description`).
#
# Запуск:  python3 scripts/build-descriptions.py [путь-к-export.xlsx]
# По умолчанию берётся свежайший ~/Downloads/export-products-*.xlsx
# ---------------------------------------------------------------------------
import sys, os, re, json, glob

try:
    import openpyxl
except ImportError:
    sys.exit('Нужен openpyxl:  pip3 install openpyxl')

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT  = os.path.join(ROOT, 'deploy', 'seed', 'category-descriptions.json')

def find_xlsx():
    if len(sys.argv) > 1:
        return sys.argv[1]
    cands = sorted(glob.glob(os.path.expanduser('~/Downloads/export-products-*.xlsx')),
                   key=os.path.getmtime, reverse=True)
    if not cands:
        sys.exit('Не найден ~/Downloads/export-products-*.xlsx — укажите файл аргументом')
    return cands[0]

ALLOWED = {'p', 'br', 'strong', 'b', 'em', 'i', 'ul', 'ol', 'li', 'a', 'h2', 'h3', 'h4'}

def clean(raw):
    if not raw:
        return ''
    s = str(raw)
    # комментарии (в т.ч. вордовские <!--[if !supportLists]-->)
    s = re.sub(r'<!--[\s\S]*?-->', '', s)
    # выкидываем целиком
    s = re.sub(r'<(iframe|script|style)[\s\S]*?</\1>', '', s, flags=re.I)
    s = re.sub(r'<img[^>]*>', '', s, flags=re.I)
    # разворачиваем контейнеры (тег убираем, содержимое оставляем)
    s = re.sub(r'</?(span|font|table|tbody|thead|tr|td|th|div|section|article)[^>]*>', '', s, flags=re.I)
    # h1 → h2
    s = re.sub(r'<(/?)h1[^>]*>', r'<\1h2>', s, flags=re.I)
    # у <a> оставляем только href, у остальных тегов — срезаем все атрибуты
    def strip_attrs(m):
        tag = m.group(1).lower()
        if tag == 'a':
            href = re.search(r'href\s*=\s*"([^"]*)"', m.group(0), flags=re.I)
            return '<a href="%s">' % href.group(1) if href else '<a>'
        return '<%s>' % tag
    s = re.sub(r'<([a-zA-Z0-9]+)(\s[^>]*)?>', strip_attrs, s)
    # сущности
    for a, b in (('&nbsp;', ' '), ('&laquo;', '«'), ('&raquo;', '»'), ('&mdash;', '—'),
                 ('&ndash;', '–'), ('&rsquo;', '’'), ('&lsquo;', '‘'),
                 ('&ldquo;', '«'), ('&rdquo;', '»'), ('&quot;', '"'), ('&amp;', '&')):
        s = s.replace(a, b)
    # незнакомые теги — вон (содержимое оставляем)
    def drop_unknown(m):
        return m.group(0) if m.group(1).lower() in ALLOWED else ''
    s = re.sub(r'</?([a-zA-Z0-9]+)\s*/?>', drop_unknown, s)
    # нормализация
    s = re.sub(r'<br\s*/?>', '<br>', s, flags=re.I)
    s = re.sub(r'(<br>\s*){2,}', '<br>', s)
    s = re.sub(r'<a[^>]*>\s*</a>', '', s, flags=re.I)
    s = re.sub(r'<(strong|b|em|i)>\s*</\1>', '', s, flags=re.I)
    s = re.sub(r'<p>\s*(<br>)?\s*</p>', '', s, flags=re.I)
    s = re.sub(r'<br>\s*</p>', '</p>', s, flags=re.I)
    s = re.sub(r'<p>\s*<br>', '<p>', s, flags=re.I)
    def _to_h3(m):
        txt = m.group(1).strip()
        return '<h3>%s</h3>' % txt if 0 < len(txt) <= 90 and not txt.endswith(('.', ':', '!', ',')) else m.group(0)
    s = re.sub(r'<p><strong>([^<]{1,120})</strong></p>', _to_h3, s, flags=re.I)
    s = re.sub(r'[ \t ]+', ' ', s)
    s = re.sub(r'>\s+<', '><', s)
    s = re.sub(r'\s{2,}', ' ', s)
    return s.strip()

def plain_len(html):
    return len(re.sub(r'<[^>]+>', '', html).strip())

def main():
    xlsx = find_xlsx()
    print('Источник:', xlsx)
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    if 'Export Groups Sheet' not in wb.sheetnames:
        sys.exit('В файле нет листа «Export Groups Sheet» (нужен экспорт Товары + Группы)')
    ws = wb['Export Groups Sheet']
    rows = list(ws.iter_rows(values_only=True))[1:]

    out = {}
    for r in rows:
        name = (r[1] or '').strip()
        if not name:
            continue
        after, before = clean(r[8]), clean(r[7])
        html = after if plain_len(after) >= 40 else (before or after)
        if plain_len(html) >= 40:
            out[name] = html

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, 'w', encoding='utf-8') as f:
        json.dump(out, f, ensure_ascii=False, indent=2, sort_keys=True)
    print('Записано %d описаний → %s' % (len(out), os.path.relpath(OUT, ROOT)))
    for k in sorted(out):
        print('  • %-55s %d симв.' % (k[:55], plain_len(out[k])))

if __name__ == '__main__':
    main()
